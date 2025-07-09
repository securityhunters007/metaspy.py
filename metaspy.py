#!/usr/bin/env python3

import argparse
import os
import json
import csv
from datetime import datetime

# Import extraction libraries
from pypdf import PdfReader
import docx
import exiftool
from pptx import Presentation
from openpyxl import load_workbook

# --- Metadata Extraction Functions ---

def extract_pdf_metadata(file_path):
    """Extracts metadata from a PDF file."""
    try:
        with open(file_path, 'rb') as f:
            reader = PdfReader(f)
            info = reader.metadata
            creation_date = info.creation_date
            mod_date = info.modification_date
            return {
                "File Type": "PDF",
                "Title": info.title,
                "Author": info.author,
                "Creator": info.creator,
                "Producer": info.producer,
                "Creation Date": creation_date.strftime("%Y-%m-%d %H:%M:%S") if creation_date else None,
                "Modification Date": mod_date.strftime("%Y-%m-%d %H:%M:%S") if mod_date else None,
            }
    except Exception as e:
        return {"Error": f"Could not process PDF: {e}"}


def extract_docx_metadata(file_path):
    """Extracts metadata from a DOCX file."""
    try:
        doc = docx.Document(file_path)
        props = doc.core_properties
        return {
            "File Type": "DOCX",
            "Author": props.author,
            "Last Modified By": props.last_modified_by,
            "Revision": props.revision,
            "Created": props.created.strftime("%Y-%m-%d %H:%M:%S") if props.created else None,
            "Modified": props.modified.strftime("%Y-%m-%d %H:%M:%S") if props.modified else None,
            "Title": props.title,
            "Subject": props.subject,
        }
    except Exception as e:
        return {"Error": f"Could not process DOCX: {e}"}

def extract_exiftool_metadata(file_path):
    """Extracts all possible metadata from an image using ExifTool."""
    try:
        with exiftool.ExifToolHelper() as et:
            metadata = et.get_metadata(file_path)[0] 
            cleaned_metadata = {key.split(':')[-1]: value for key, value in metadata.items()}
            return cleaned_metadata
    except Exception as e:
        return {"Error": f"Could not process image with ExifTool: {e}"}

def extract_pptx_metadata(file_path):
    """Extracts metadata from a PPTX file."""
    try:
        prs = Presentation(file_path)
        props = prs.core_properties
        return {
            "File Type": "PPTX",
            "Author": props.author,
            "Last Modified By": props.last_modified_by,
            "Revision": props.revision,
            "Created": props.created.strftime("%Y-%m-%d %H:%M:%S") if props.created else None,
            "Modified": props.modified.strftime("%Y-%m-%d %H:%M:%S") if props.modified else None,
            "Title": props.title,
            "Subject": props.subject,
        }
    except Exception as e:
        return {"Error": f"Could not process PPTX: {e}"}

def extract_xlsx_metadata(file_path):
    """Extracts metadata from an XLSX file."""
    try:
        wb = load_workbook(file_path)
        props = wb.properties
        return {
            "File Type": "XLSX",
            "Creator": props.creator,
            "Last Modified By": props.lastModifiedBy,
            "Created": props.created.strftime("%Y-%m-%d %H:%M:%S") if props.created else None,
            "Modified": props.modified.strftime("%Y-%m-%d %H:%M:%S") if props.modified else None,
            "Title": props.title,
            "Subject": props.subject,
        }
    except Exception as e:
        return {"Error": f"Could not process XLSX: {e}"}

# --- Output Generation Functions (Corrected) ---

def save_as_txt(data, filename):
    """Saves extracted metadata to a TXT file."""
    with open(filename, 'w', encoding='utf-8') as f:
        for item in data:
            f.write(f"--- Metadata for: {item['file']} ---\n")
            for key, value in item['metadata'].items():
                f.write(f"{key}: {value}\n")
            if 'Geolocation' in item:
                f.write(f"Geolocation: {item['Geolocation']}\n")
            f.write("\n")
    print(f"✅ Report saved to {filename}")

def save_as_json(data, filename):
    """Saves extracted metadata to a JSON file."""
    with open(filename, 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=4)
    print(f"✅ Report saved to {filename}")

def save_as_csv(data, filename):
    """Saves extracted metadata to a CSV file."""
    with open(filename, 'w', newline='', encoding='utf-8') as f:
        flat_data = []
        for item in data:
            base_info = {'File': item['file']}
            if 'Geolocation' in item:
                base_info['Geolocation'] = item['Geolocation']
            base_info.update(item['metadata'])
            flat_data.append(base_info)
        
        if not flat_data:
            print("⚠️ No data to write.")
            return

        headers = sorted(list(set(key for d in flat_data for key in d.keys())))
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(flat_data)
    print(f"✅ Report saved to {filename}")


# --- Main Application Logic ---

def main():
    """Main function to parse arguments and orchestrate extraction."""
    parser = argparse.ArgumentParser(
        description="MetaSpy: A metadata extraction tool for various file types.",
        epilog="Example: python metaspy.py mydoc.docx myphoto.jpg mydata.xlsx -o json"
    )
    parser.add_argument("files", nargs="+", help="One or more file paths to analyze.")
    parser.add_argument(
        "--output", "-o", 
        choices=["txt", "csv", "json", "print"], 
        default="print", 
        help="The format for the output report (default: print to console)."
    )
    args = parser.parse_args()

    file_handlers = {
        '.pdf': extract_pdf_metadata,
        '.docx': extract_docx_metadata,
        '.jpg': extract_exiftool_metadata,
        '.jpeg': extract_exiftool_metadata,
        '.png': extract_exiftool_metadata,
        '.tiff': extract_exiftool_metadata,
        '.gif': extract_exiftool_metadata,
        '.bmp': extract_exiftool_metadata,
        '.pptx': extract_pptx_metadata,
        '.xlsx': extract_xlsx_metadata,
    }

    print("🕵️  Starting MetaSpy analysis (v1.3 with Office Suite Support)...")
    all_metadata = []

    for file_path in args.files:
        if not os.path.exists(file_path):
            print(f"❌ Error: File not found at '{file_path}'")
            continue

        file_ext = os.path.splitext(file_path)[1].lower()
        handler = file_handlers.get(file_ext)

        if handler:
            print(f"📄 Analyzing {file_path}...")
            metadata = handler(file_path)
            item = {"file": file_path, "metadata": metadata}
            
            if 'GPSLatitude' in metadata and 'GPSLongitude' in metadata:
                try:
                    lat, lon = metadata['GPSLatitude'], metadata['GPSLongitude']
                    maps_link = f"https://www.google.com/maps/search/?api=1&query={lat},{lon}"
                    item['Geolocation'] = maps_link
                except Exception:
                    pass
            all_metadata.append(item)
        else:
            print(f"⚠️ Warning: Unsupported file type for '{file_path}'. Skipping.")

    if args.output == "print":
        for item in all_metadata:
            print(f"\n--- Metadata for: {item['file']} ---")
            if 'Error' in item['metadata']:
                print(f"Error: {item['metadata']['Error']}")
            else:
                for key, value in item['metadata'].items():
                    print(f"  {key}: {value}")
                if 'Geolocation' in item:
                    print(f"  📍 Geolocation Link: {item['Geolocation']}")
        print("\n✅ Analysis complete.")
    else:
        output_handlers = { 'txt': save_as_txt, 'json': save_as_json, 'csv': save_as_csv }
        output_function = output_handlers.get(args.output)
        if output_function:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_filename = f"metaspy_report_{timestamp}.{args.output}"
            output_function(all_metadata, output_filename)

if __name__ == "__main__":
    main()